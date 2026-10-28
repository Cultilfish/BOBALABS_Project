"""Excel report of what clients ask about: the distribution of question
types across the dataset, the train/test split, and which examples each
training-selection strategy picks (even spread vs category-balanced).

Usage: python scripts/question_distribution.py [--n 15] [--out docs/question_distribution.xlsx]
Requires openpyxl (pip install -r requirements-dev.txt).
"""
from __future__ import annotations

import argparse
from collections import Counter

from common import ROOT, DATA, TEST_IDS, spread
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.categories import (QUESTION_CATEGORIES, STATEMENT_CATEGORIES, all_categories,
                            balanced_sample, categorize)
from app.conversation import build_examples, load_conversations, split_by_contact

FONT = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="534AB7")
TOTAL_FILL = PatternFill("solid", fgColor="EEEDFE")
THIN = Side(style="thin", color="D3D1C7")
BORDER = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)


def style_table(ws, n_rows: int, n_cols: int, widths: list[int], wrap_cols=()):
    for r in range(1, n_rows + 1):
        for c in range(1, n_cols + 1):
            cell = ws.cell(r, c)
            cell.font = Font(name=FONT, size=10, bold=(r == 1), color="FFFFFF" if r == 1 else "000000")
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=(r == 1 or c in wrap_cols))
            if r == 1:
                cell.fill = HEAD_FILL
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--n", type=int, default=15, help="training examples picked per strategy")
    ap.add_argument("--out", default=str(ROOT / "docs" / "question_distribution.xlsx"))
    args = ap.parse_args()

    examples = build_examples(load_conversations(DATA))
    train, _ = split_by_contact(examples, TEST_IDS)  # same objects as `examples`
    spread_ids = {id(e) for e in spread(train, args.n)}
    balanced_ids = {id(e) for e in balanced_sample(train, args.n)}
    counts = Counter(categorize(e.client_sequence) for e in examples)
    cats = sorted(all_categories(), key=lambda c: (-counts[c[0]], c[0]))
    cats = [c for c in cats if counts[c[0]] > 0]

    wb = Workbook()

    # ---- Questions (raw data; Summary formulas count from here) -------------
    q = wb.active
    q.title = "Questions"
    q.append(["#", "Conversation", "Scenario", "Split", "Question type",
              "Picked: even spread", "Picked: balanced", "Client message(s)", "Real consultant reply"])
    for i, e in enumerate(examples, 1):
        q.append([i, e.contact_id, e.scenario, "Test" if e.contact_id in TEST_IDS else "Train",
                  categorize(e.client_sequence),
                  "Yes" if id(e) in spread_ids else "No",
                  "Yes" if id(e) in balanced_ids else "No",
                  "\n".join(e.client_sequence), e.consultant_reply])
    last = len(examples) + 1
    style_table(q, last, 9, [5, 12, 30, 7, 26, 10, 10, 55, 70], wrap_cols=(3, 8, 9))
    q.auto_filter.ref = f"A1:I{last}"

    # ---- Summary ------------------------------------------------------------
    s = wb.create_sheet("Summary", 0)
    s.append(["Question type", "What it covers", "All examples", "Share of all",
              "Train", "Test", f"Picked: even spread (n={args.n})", f"Picked: balanced (n={args.n})"])
    rng = lambda col: f"Questions!${col}$2:${col}${last}"  # noqa: E731
    for r, (name, desc) in enumerate(cats, 2):
        s.append([
            name, desc,
            f'=COUNTIF({rng("E")},A{r})',
            f"=IF($C${len(cats) + 2}=0,0,C{r}/$C${len(cats) + 2})",
            f'=COUNTIFS({rng("E")},A{r},{rng("D")},"Train")',
            f'=COUNTIFS({rng("E")},A{r},{rng("D")},"Test")',
            f'=COUNTIFS({rng("E")},A{r},{rng("F")},"Yes")',
            f'=COUNTIFS({rng("E")},A{r},{rng("G")},"Yes")',
        ])
    tot = len(cats) + 2
    s.append(["Total", ""] + [f"=SUM({c}2:{c}{tot - 1})" for c in "CDEFGH"])
    s.append(["Question types covered", "", "", "", "", "",
              f'=COUNTIF(G2:G{tot - 1},">0")', f'=COUNTIF(H2:H{tot - 1},">0")'])
    style_table(s, tot + 1, 8, [30, 52, 12, 11, 8, 8, 16, 16], wrap_cols=(2, 7, 8))
    for r in range(2, tot + 2):
        s.cell(r, 4).number_format = "0.0%"
    for r in (tot, tot + 1):
        for c in range(1, 9):
            s.cell(r, c).font = Font(name=FONT, size=10, bold=True)
            s.cell(r, c).fill = TOTAL_FILL

    note_row = tot + 3
    notes = [
        "How to read this",
        "One row per client turn: the client's message(s) that a real consultant then replied to "
        f"({len(examples)} turns from {len({e.contact_id for e in examples})} conversations).",
        "Question types come from keyword rules in app/categories.py (see the 'Type rules' sheet). "
        "They are deterministic, so the same message always gets the same type.",
        f"Test = conversations held out from training ({', '.join(sorted(TEST_IDS))}).",
        "The last two columns compare two ways of choosing training examples. Even spread "
        "follows the data's natural mix. Balanced takes examples round-robin across types, rarest "
        "first, so rare situations (rejections, interviews) are included, in the same spirit as "
        "oversampling rare classes in image detection, without duplicating any example.",
    ]
    for i, text in enumerate(notes):
        c = s.cell(note_row + i, 1, text)
        c.font = Font(name=FONT, size=10, bold=(i == 0))
        s.merge_cells(start_row=note_row + i, start_column=1, end_row=note_row + i, end_column=8)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        s.row_dimensions[note_row + i].height = 15 if i == 0 else 30

    chart = BarChart()
    chart.type = "bar"
    chart.title = "What clients ask about (all examples)"
    chart.y_axis.title = "Client turns"
    chart.legend = None
    chart.add_data(Reference(s, min_col=3, min_row=1, max_row=tot - 1), titles_from_data=True)
    chart.set_categories(Reference(s, min_col=1, min_row=2, max_row=tot - 1))
    chart.x_axis.scaling.orientation = "maxMin"
    chart.height, chart.width = 9, 16
    s.add_chart(chart, "J2")

    chart2 = BarChart()
    chart2.type = "bar"
    chart2.title = f"Training picks by type: even spread vs balanced (n={args.n})"
    chart2.add_data(Reference(s, min_col=7, max_col=8, min_row=1, max_row=tot - 1), titles_from_data=True)
    chart2.set_categories(Reference(s, min_col=1, min_row=2, max_row=tot - 1))
    chart2.x_axis.scaling.orientation = "maxMin"
    chart2.height, chart2.width = 9, 16
    s.add_chart(chart2, "J22")

    # ---- Type rules -----------------------------------------------------------
    t = wb.create_sheet("Type rules")
    t.append(["Order", "Question type", "Applies to", "Keyword rule (regex, first match wins)"])
    for i, (name, _, pat) in enumerate(QUESTION_CATEGORIES, 1):
        t.append([i, name, "Messages with a '?'", pat])
    for i, (name, _, pat) in enumerate(STATEMENT_CATEGORIES, len(QUESTION_CATEGORIES) + 1):
        t.append([i, name, "Messages without a '?'", pat])
    style_table(t, t.max_row, 4, [7, 30, 22, 90], wrap_cols=(4,))

    wb.save(args.out)
    print(f"Wrote {args.out}")
    print("Open it in Excel (formulas calculate on open), or run LibreOffice recalculation first.")


if __name__ == "__main__":
    main()
